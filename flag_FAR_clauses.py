import pandas as pd
import re
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

# Main Program
def annotate_contract(xls_in, contract_in, save_path):
    # Read FAR Matrix and clean it
    xls_matrix_path = xls_in
    df_clean = read_far_matrix(xls_matrix_path)
    xls_clauses = df_clean['Clause'].tolist()
    acceptance_status = df_clean['Acceptance Status*'].tolist()

    # Read Contract Text
    contract_text = read_docx(contract_in)

    # Flag Clauses in Contract Text
    flagged_clauses = flag_clauses(contract_text, xls_clauses, acceptance_status)

    # Create New Annotated Document
    new_doc = Document()
    original_doc = Document(contract_in)

    add_color_code(new_doc)
    for paragraph in original_doc.paragraphs:
        annotate_paragraph(paragraph, flagged_clauses, new_doc)
    
    # Save Annotated Document
    new_doc.save(save_path)

# Acceptance statuses from least to most restrictive, and the highlight used for each
STATUS_SEVERITY = {'OK': 0, 'C': 1, 'REMOVE': 2}
STATUS_COLORS = {
    'OK': WD_COLOR_INDEX.BRIGHT_GREEN,   # Green for OK
    'C': WD_COLOR_INDEX.TURQUOISE,       # Turquoise for Conditional (would be blue but it's hard to read)
    'REMOVE': WD_COLOR_INDEX.RED,        # Red for Remove
}
# A FAR clause number, such as 52.204-21
_CLAUSE_NUMBER = re.compile(r'^\d{2}\.\d{3}-\d{1,4}$')

# Function to read FAR Matrix from Excel and return clean DataFrame.
# Keeps only rows whose Clause is a real clause number: the matrix also has 326 archived rows whose
# Clause is "OLD" and one "RESERVED" row, which used to match ordinary words like "hold" and "reserved".
# Statuses are normalized ("Remove", " C", "OK " were never highlighted), and a clause listed more than
# once keeps its most restrictive status, so the tool never shows a clause as more acceptable than OSP does.
def read_far_matrix(excel_path):
    df = pd.read_excel(excel_path)
    df_clean = df.dropna(subset=['Clause ', 'Acceptance Status*'])
    df_clean = df_clean[['Clause ', 'Acceptance Status*']]
    df_clean = df_clean.rename(columns={'Clause ': 'Clause'})
    df_clean['Clause'] = df_clean['Clause'].astype(str).str.strip()
    df_clean['Acceptance Status*'] = df_clean['Acceptance Status*'].astype(str).str.strip().str.upper()
    df_clean = df_clean[df_clean['Clause'].str.match(_CLAUSE_NUMBER) & df_clean['Acceptance Status*'].isin(STATUS_SEVERITY)]
    df_clean = (df_clean.assign(severity=df_clean['Acceptance Status*'].map(STATUS_SEVERITY))
                .sort_values('severity', kind='stable')
                .drop_duplicates('Clause', keep='last')
                .drop(columns='severity')
                .sort_index())
    return df_clean

# Matches a clause number on its own, so "52.204-2" does not match inside "52.204-21"
def _clause_pattern(clause):
    return re.compile(rf'(?<![\d.]){re.escape(clause)}(?!\d)')

# Returns the more restrictive of two statuses
def _stricter(status_a, status_b):
    return max(status_a, status_b, key=lambda status: STATUS_SEVERITY.get(status, -1))

# Function to read text from a .docx file
def read_docx(file_path):
    doc = Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    return '\n'.join(full_text)

# Function to flag clauses in contract text
def flag_clauses(contract_text, xls_clauses, acceptance_status):
    flagged_clauses = {}
    for clause, status in zip(xls_clauses, acceptance_status):
        if _clause_pattern(clause).search(contract_text):
            flagged_clauses[clause] = _stricter(flagged_clauses.get(clause, status), status)
    return flagged_clauses

def add_color_code(new_doc):
    paragraph = new_doc.add_paragraph()

    lines = ['FAR Acceptance Status Color Code', 'Green = OK', 'Turquoise = Conditional', 'Red = Remove']
    colors = [WD_COLOR_INDEX.WHITE, WD_COLOR_INDEX.BRIGHT_GREEN, WD_COLOR_INDEX.TURQUOISE, WD_COLOR_INDEX.RED]

    for line, color in zip(lines, colors):
        run = paragraph.add_run(line)
        if lines[0] == line: 
            run.bold = True 
            run.underline = True
        run.font.highlight_color = color
        run.add_break()


# Function to annotate a paragraph based on flagged clauses and their acceptance status
def annotate_paragraph(paragraph, flagged_clauses, new_doc):
    new_paragraph = new_doc.add_paragraph()
    for run in paragraph.runs:
        new_run = new_paragraph.add_run(run.text, run.style.name)
        new_run.font.highlight_color = run.font.highlight_color if run.font.highlight_color else None
        # When a run mentions several clauses, show the most restrictive status among them
        strictest = None
        for clause, status in flagged_clauses.items():
            if _clause_pattern(clause).search(run.text):
                strictest = status if strictest is None else _stricter(strictest, status)
        if strictest in STATUS_COLORS:
            new_run.font.highlight_color = STATUS_COLORS[strictest]