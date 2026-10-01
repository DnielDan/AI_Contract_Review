"""Tests for FAR clause detection and highlighting."""
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

import main
from flag_FAR_clauses import annotate_contract, flag_clauses, read_far_matrix

# Clauses whose numbers do not contain any other clause number, one per acceptance status
EXPECTED_HIGHLIGHTS = {
    "52.203-18": WD_COLOR_INDEX.BRIGHT_GREEN,  # OK
    "52.202-1": WD_COLOR_INDEX.TURQUOISE,      # C (conditional)
    "52.203-16": WD_COLOR_INDEX.RED,           # REMOVE
}


def test_flag_clauses_finds_listed_clauses_only():
    text = "This agreement incorporates FAR 52.203-18 and 52.203-16 by reference."
    flagged = flag_clauses(text, ["52.203-18", "52.203-16", "52.202-1"], ["OK", "REMOVE", "C"])
    assert flagged == {"52.203-18": "OK", "52.203-16": "REMOVE"}


def test_annotated_document_highlights_clauses_by_status(tmp_path):
    contract = Document()
    contract.add_paragraph("This paragraph mentions no FAR clauses.")
    for clause in EXPECTED_HIGHLIGHTS:
        contract.add_paragraph(f"The Contractor shall comply with FAR {clause}.")
    contract_path = tmp_path / "contract.docx"
    output_path = tmp_path / "annotated.docx"
    contract.save(contract_path)

    annotate_contract(main.FAR_CLAUSE_MATRIX_PATH, str(contract_path), str(output_path))

    highlights = {}
    for paragraph in Document(output_path).paragraphs:
        for run in paragraph.runs:
            highlights[run.text.strip()] = run.font.highlight_color
    assert "FAR Acceptance Status Color Code" in highlights
    assert highlights["This paragraph mentions no FAR clauses."] is None
    for clause, color in EXPECTED_HIGHLIGHTS.items():
        assert highlights[f"The Contractor shall comply with FAR {clause}."] == color


def _annotate(tmp_path, paragraphs):
    contract = Document()
    for text in paragraphs:
        contract.add_paragraph(text)
    contract.save(tmp_path / "contract.docx")
    annotate_contract(main.FAR_CLAUSE_MATRIX_PATH, str(tmp_path / "contract.docx"), str(tmp_path / "annotated.docx"))
    return {run.text.strip(): run.font.highlight_color
            for paragraph in Document(tmp_path / "annotated.docx").paragraphs for run in paragraph.runs}


def test_archived_and_reserved_rows_do_not_match_ordinary_words(tmp_path):
    # The matrix has 326 archived rows whose Clause is "OLD" and one "RESERVED" row; they used to
    # highlight any text containing "hold", "withhold", "threshold", or "reserved" (often as green "OK")
    far = read_far_matrix(main.FAR_CLAUSE_MATRIX_PATH)
    assert not far["Clause"].isin(["OLD", "RESERVED"]).any()
    text = "The Subrecipient shall indemnify and hold harmless the Endowment. All rights reserved."
    assert flag_clauses(text, far["Clause"].tolist(), far["Acceptance Status*"].tolist()) == {}
    assert _annotate(tmp_path, [text])[text] is None


def test_mistyped_statuses_are_highlighted(tmp_path):
    # Raw statuses in the matrix: "Remove", " C", "OK ", "REMOVE " were never highlighted before
    expected = {
        "52.216-2": WD_COLOR_INDEX.RED,          # "Remove"
        "52.215-20": WD_COLOR_INDEX.TURQUOISE,   # " C"
        "52.232-12": WD_COLOR_INDEX.BRIGHT_GREEN,  # "OK "
        "52.246-8": WD_COLOR_INDEX.RED,          # "REMOVE "
    }
    highlights = _annotate(tmp_path, [f"See FAR {clause}." for clause in expected])
    for clause, color in expected.items():
        assert highlights[f"See FAR {clause}."] == color


def test_clause_numbers_match_exactly():
    # "52.216-3" is a different clause from "52.216-30", and both are in the matrix
    assert flag_clauses("Clause 52.216-30 applies.", ["52.216-3", "52.216-30"], ["REMOVE", "REMOVE"]) == {"52.216-30": "REMOVE"}
    assert flag_clauses("FAR 52.216-3(a) applies.", ["52.216-3", "52.216-30"], ["REMOVE", "REMOVE"]) == {"52.216-3": "REMOVE"}


def test_conflicting_statuses_use_the_most_restrictive(tmp_path):
    # 52.223-2 is listed twice in the matrix, as C and as REMOVE (the second row's title belongs to 52.225-2)
    far = read_far_matrix(main.FAR_CLAUSE_MATRIX_PATH)
    assert far.loc[far["Clause"] == "52.223-2", "Acceptance Status*"].tolist() == ["REMOVE"]
    text = "FAR 52.203-18 and FAR 52.203-16 are incorporated."  # OK and REMOVE in one paragraph
    assert _annotate(tmp_path, [text])[text] == WD_COLOR_INDEX.RED
