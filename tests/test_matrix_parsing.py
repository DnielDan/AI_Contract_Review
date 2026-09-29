"""Tests that the FAR and Ts&Cs matrices load into the structures the pipeline expects."""
import openpyxl

import main
from flag_FAR_clauses import read_far_matrix
from flag_problem_language import extract_sheet_data

SKIPPED_SHEETS = {"INDEX", "template", "CONTACTS"}


def test_far_matrix_has_clauses_and_statuses():
    far = read_far_matrix(main.FAR_CLAUSE_MATRIX_PATH)
    assert list(far.columns) == ["Clause", "Acceptance Status*"]
    assert len(far) > 700
    statuses = {str(status).strip().upper() for status in far["Acceptance Status*"]}
    assert {"OK", "C", "REMOVE"} <= statuses


def test_tnc_matrix_sheets_parse_with_aligned_columns():
    workbook = openpyxl.load_workbook(main.TNC_MATRIX_PATH)
    sheets_with_problems = 0
    for name in workbook.sheetnames:
        if name in SKIPPED_SHEETS:
            continue
        data = extract_sheet_data(workbook[name])
        # "Why" and "1st response" are read row by row alongside each common problem
        assert len(data["Why"]) == len(data["Common Problems"])
        assert len(data["1st response to Sponsor"]) == len(data["Common Problems"])
        sheets_with_problems += bool(data["Common Problems"])
    assert sheets_with_problems >= 20


def test_tnc_matrix_is_the_current_version():
    # The March 2024 update added this Intellectual Property problem; an older copy lacks it
    workbook = openpyxl.load_workbook(main.TNC_MATRIX_PATH)
    problems = extract_sheet_data(workbook["Intellectual Property"])["Common Problems"]
    assert any("split the costs of patent preparation" in str(problem) for problem in problems)
